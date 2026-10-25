from docx import Document
d = Document(r"E:\deepseek\deliverables\conv-20260928-3b54e160\local-ai-toolset-guide-v1.docx")
text = "\n".join(p.text for p in d.paragraphs)
for t in d.tables:
    for row in t.rows:
        text += "\n" + " | ".join(c.text for c in row.cells)
print("  段落数: %d, 表格数: %d, 总字符: %d" % (len(d.paragraphs), len(d.tables), len(text)))
for kw in ["本地 AI 工具集", "FLUX.1-schnell", "模型变更后的运维规则", "版本 1.4", "vision_status", "生图的时间预算", "硬件寿命"]:
    print("  含 %-24s %s" % (kw, "是" if kw in text else "否"))