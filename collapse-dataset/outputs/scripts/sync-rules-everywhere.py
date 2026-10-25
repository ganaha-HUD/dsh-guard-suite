import json
import os
import re

HOME = os.path.expanduser("~")
W = r"E:\Deepseek\AGENTS.md"
G = os.path.join(HOME, ".dsh", "AGENTS.md")
GSRC = r"E:\Deepseek\policy\deploy\global-agents.md"

w = open(W, encoding="utf-8").read()
g = open(G, encoding="utf-8").read()
gsrc = open(GSRC, encoding="utf-8").read()

print("=== 4) 关键规则条目在两份 AGENTS 中的覆盖情况 ===")
CHECKS = [
    ("强制中文思考", "思考语言"),
    ("AI 工具：文档走文字层/OCR", "vision_read_document"),
    ("AI 工具：FLUX 生图", "FLUX.1-schnell"),
    ("生图预算", "生图预算"),
    ("运维：换模型后重启自检", "更换模型后必须重启服务端"),
    ("机器寿命规则", "兼顾机器寿命"),
    ("NPU 1024 上限", "1024 token"),
    ("VLM 幻觉提示要求", "未与文字层/OCR 交叉校验"),
    ("reference_file 只读", "reference_file"),
    ("AI 文件放 D:\\AI_file", "D:\\AI_file"),
    ("温度监控命令", "Thermal Zone Information"),
]
print("  %-34s %-10s %-10s" % ("检查项", "工作区", "全局"))
for label, kw in CHECKS:
    print("  %-34s %-10s %-10s" % (label, "有" if kw in w else "缺!", "有" if kw in g else "缺!"))

print()
print("=== 5) 全局副本与源文件是否一致 ===")
print("  global-agents.md 与 ~/.dsh/AGENTS.md 内容相同:", gsrc == g)

print()
print("=== 6) 潜在矛盾：'不要频繁重启服务端' vs '换模型后必须重启' ===")
for name, txt in (("工作区", w), ("全局", g)):
    for kw in ("不要为预热空跑或频繁重启服务端", "更换模型后必须重启服务端"):
        for m in re.finditer(re.escape(kw), txt):
            seg = txt[max(0, m.start() - 60):m.start() + 140].replace("\n", " ")
            print("  [%s] %s" % (name, seg.strip()[:200]))
            break

print()
print("=== 7) DOCX 是否已包含最新规则（§12 运维 + 中文思考 + FLUX） ===")
try:
    from docx import Document
    d = Document(r"E:\Deepseek\deliverables\conv-20260928-3b54e160\local-ai-toolset-guide-v1.docx")
    text = "\n".join(p.text for p in d.paragraphs)
    for t in d.tables:
        for row in t.rows:
            text += "\n" + " | ".join(c.text for c in row.cells)
    for label, kw in (("中文思考", "一律使用中文"), ("FLUX 生图", "FLUX.1-schnell"),
                      ("生图预算", "生图的时间预算"), ("运维规则 §12", "更换模型后必须重启"),
                      ("健康检查说明", "vision_status"), ("寿命规则", "硬件寿命")):
        print("  %-20s %s" % (label, "有" if kw in text else "缺 ← 需补"))
except Exception as exc:
    print("  读取 DOCX 失败:", exc)

print()
print("=== 8) MCP 注册与对话登记 ===")
for prof in ("desktop", "web"):
    p = os.path.join(HOME, ".dsh", "profiles", prof, "cordis.patch.yml")
    if os.path.isfile(p):
        t = open(p, encoding="utf-8").read()
        print("  %-8s mcp-vision 已注册: %s   insert 包裹: %s" % (prof, "mcp-vision" in t, "- insert:" in t))
    else:
        print("  %-8s 配置文件缺失" % prof)
cr = json.load(open(r"E:\Deepseek\policy\conversation-registry.json", encoding="utf-8"))
txt = json.dumps(cr, ensure_ascii=False)
print("  对话登记含本对话 3b54e160:", "3b54e160" in txt)
print("  登记条目数:", len(cr) if isinstance(cr, list) else len(cr.get("conversations", cr)))